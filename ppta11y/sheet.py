"""The alt-text spreadsheet: read rows; record results in new columns of a *copy*."""
from __future__ import annotations

import shutil
from copy import copy
from pathlib import Path

import openpyxl

from .com import ToolError
from .session import same_file

MARKER = "ppt-a11y results copy"  # stored in the workbook's Subject property
LOG_MARKER = "ppt-a11y change log"  # a workbook holding only the change log (never a results copy)


def _put(ws, row: int, col: int, v):
    """Write a cell safely: PowerPoint line breaks become ' / ', characters Excel cannot store are dropped,
    and text starting with '=' stays text (never a formula)."""
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    if isinstance(v, str):
        v = ILLEGAL_CHARACTERS_RE.sub("", v.replace("\x0b", " / ").replace("\r", " / "))
    c = ws.cell(row=row, column=col, value=v)
    if isinstance(v, str) and v.startswith("="):
        c.data_type = "s"
    return c


def read(path: str, sheet: str | None = None) -> dict:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    headers = [str(c.value).strip() if c.value is not None else f"col{c.column}" for c in ws[1]]
    rows = []
    for r in ws.iter_rows(min_row=2):
        vals = [c.value for c in r]
        if all(v is None for v in vals):
            continue
        rows.append({"row": r[0].row, **{h: v for h, v in zip(headers, vals) if v is not None}})
    return {"path": str(path), "sheet": ws.title, "headers": headers, "rows": rows,
            "other_sheets": [w.title for w in wb.worksheets if w.title != ws.title]}


def _subject(path: Path) -> str:
    try:
        return openpyxl.load_workbook(path, read_only=True).properties.subject or ""
    except Exception:
        return ""


def _is_results_copy(path: Path) -> bool:
    return _subject(path).startswith(MARKER)


def record(source: str, dest: str, updates: dict, sheet: str | None = None) -> dict:
    """Write `updates` ({row: {header: value}}) into `dest`, a copy of `source` made on first use.
    New headers become new columns at the right. Every column that exists in `source` is read-only, and
    `dest` must be a results copy made by this function (so swapped arguments cannot touch the original)."""
    src, out = Path(source).resolve(), Path(dest).resolve()
    if not src.exists():
        raise ToolError(f"source spreadsheet not found: {src}")
    if same_file(src, out):
        raise ToolError("write results to a copy, not the original spreadsheet")
    if _is_results_copy(src):
        raise ToolError(f"{src} is a results copy; pass the ORIGINAL spreadsheet as source")
    if out.exists() and not _is_results_copy(out):
        raise ToolError(f"{out} exists and is not a results copy made by this tool; refusing to modify it")
    created = not out.exists()
    if created:
        shutil.copy2(src, out)
    wb = openpyxl.load_workbook(out)
    if created:
        wb.properties.subject = f"{MARKER} of {src.name}"
    ws = wb[sheet] if sheet else wb.worksheets[0]
    src_ws = openpyxl.load_workbook(src, read_only=True)
    src_ws = src_ws[sheet] if sheet else src_ws.worksheets[0]
    protected = {str(c.value).strip().lower() for c in next(src_ws.iter_rows(min_row=1, max_row=1))
                 if c.value is not None}
    headers = {str(c.value).strip(): c.column for c in ws[1] if c.value is not None}
    written = 0
    for row, cols in updates.items():
        row = int(row)
        if row < 2:
            raise ToolError("row numbers start at 2 (row 1 is the header)")
        for h, v in cols.items():
            if h.strip().lower() in protected:
                raise ToolError(f"refusing to edit input column '{h}' (it exists in the original spreadsheet)")
            col = headers.get(h.strip())
            if col is None:
                col = ws.max_column + 1
                cell = ws.cell(row=1, column=col, value=h.strip())
                ref = ws.cell(row=1, column=col - 1)
                cell.font, cell.fill, cell.alignment, cell.border = (copy(ref.font), copy(ref.fill),
                                                                     copy(ref.alignment), copy(ref.border))
                ws.column_dimensions[cell.column_letter].width = 45
                headers[h.strip()] = col
            c = _put(ws, row, col, v)
            c.alignment = copy(ws.cell(row=row, column=3).alignment)
            written += 1
    wb.save(out)
    return {"path": str(out), "cells_written": written, "columns": list(headers)}


CHANGE_COLUMNS = ("Slide", "Category", "WCAG", "What changed", "Kind", "Before", "After", "Verified")


def record_changes(dest: str, changes: list[dict], sheet: str = "WCAG changes", replace: bool = True) -> dict:
    """Write a change log (one row per change: Slide, Category, WCAG, What changed, Kind = invisible /
    visible, Before, After, Verified) to its own sheet of the results copy `dest` (or a new change-log
    workbook made by this tool). The input sheets are never touched: only a sheet that already is a change
    log (its first row is the change-log header) is rewritten. replace=True (default) rewrites that sheet
    on every call, so pass every change at once; replace=False appends."""
    from openpyxl.styles import Alignment, Font, PatternFill
    if not isinstance(changes, list) or not all(isinstance(c, dict) for c in changes):
        raise ToolError("changes must be a list of objects like {\"Slide\": 4, \"Category\": \"Contrast\", ...}")
    out = Path(dest).resolve()
    if out.exists():
        if not (_is_results_copy(out) or _subject(out) == LOG_MARKER):
            raise ToolError(f"{out} is not a results copy or change log made by this tool; refusing to modify it")
        wb = openpyxl.load_workbook(out)
    else:
        wb = openpyxl.Workbook()
        wb.active.title = sheet
        wb.properties.subject = LOG_MARKER
    if sheet in wb.sheetnames:
        head = tuple(wb[sheet].cell(1, i).value for i in range(1, len(CHANGE_COLUMNS) + 1))
        empty = wb[sheet].max_row <= 1 and wb[sheet].cell(1, 1).value is None
        if head != CHANGE_COLUMNS and not empty:
            raise ToolError(f"sheet '{sheet}' in {out.name} is not a change log (it may be the copied input sheet); "
                            "pass another sheet name")
        if replace:
            del wb[sheet]
    ws = wb[sheet] if sheet in wb.sheetnames else wb.create_sheet(sheet)
    if ws.max_row <= 1 and ws.cell(1, 1).value is None:
        for i, h in enumerate(CHANGE_COLUMNS, 1):
            c = ws.cell(row=1, column=i, value=h)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="861F41")
        for col, w in zip("ABCDEFGH", (7, 16, 10, 70, 11, 32, 32, 40)):
            ws.column_dimensions[col].width = w
        ws.freeze_panes = "A2"
    for ch in changes:
        r = ws.max_row + 1
        for i, h in enumerate(CHANGE_COLUMNS, 1):
            c = _put(ws, r, i, ch.get(h, ch.get(h.lower(), "")))
            c.alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(out)
    return {"path": str(out), "sheet": sheet, "rows": len(changes)}
